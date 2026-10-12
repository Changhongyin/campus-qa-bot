#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
中国计量大学校园问答机器人 —— 实验报告自动生成脚本 (python-docx)
================================================================
输入模板：templates/report_template.docx（由桌面"实验报告格式.doc"经 LibreOffice 转换而来）
输出文件：~/Desktop/实验报告-完成版.docx

功能：
  1. 填入"实验名称"；
  2. 在"实验目标""实验环境""实验步骤"标题下自动写入内容；
  3. 自动把 screenshots/ 下由 Playwright 生成的截图按比例（宽 5.5 英寸）插入到"实验步骤"对应位置；
  4. 另存为桌面"实验报告-完成版.docx"。
"""
import json
import pathlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

ROOT = pathlib.Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "templates" / "report_template.docx"
SHOTS = ROOT / "screenshots"
OUT = pathlib.Path.home() / "Desktop" / "实验报告-完成版.docx"
FONT = "宋体"
IMG_WIDTH = Inches(5.5)   # 正文可用宽度约 5.77 英寸，留出边距


def set_run_font(run, name=FONT, size=12.0, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def insert_paragraph_after(anchor, text="", bold=False, size=12.0, align=None):
    """在 anchor 段落之后插入一个新段落，返回新段落对象。"""
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if align is not None:
        para.alignment = align
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(4)
    if text:
        set_run_font(para.add_run(text), size=size, bold=bold)
    return para


def insert_image_after(anchor, image_path, caption, size=12.0):
    """在 anchor 之后插入一张居中的图片和一行图注，返回图注段落。"""
    pic_p = OxmlElement("w:p")
    anchor._p.addnext(pic_p)
    pic_para = Paragraph(pic_p, anchor._parent)
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_para.add_run().add_picture(str(image_path), width=IMG_WIDTH)
    cap = insert_paragraph_after(pic_para, caption, size=10.5,
                                 align=WD_ALIGN_PARAGRAPH.CENTER)
    return cap


def find_paragraph(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def load_answers():
    f = SHOTS / "answers.json"
    if f.exists():
        return json.loads(f.read_text(encoding="utf-8"))
    return {}


# --------------------------- 报告正文内容 ---------------------------
GOALS = [
    "1）用 AI 开发一个智能体程序，完整体验“需求分析 → 环境搭建 → 编码 → 测试 → 交付”的智能体开发流程。",
    "2）在 DeepSeek 官网注册账号、充值并申请 API Key，掌握把大模型 API 接入自己程序的方法。",
    "3）自己手写系统提示词，赋予智能体独一无二的角色——中国计量大学校园问答机器人，使其拥有本地知识库，能回答外网鲜为人知的校园问题。",
    "4）交付 AI 开发的代码级作品，并明确说明提示词的修改位置，做到“AI 写代码、自己写提示词”。",
]

ENVIRONMENT = [
    "硬件环境：MacBook（Apple Silicon，macOS）。",
    "运行环境：Node.js v26.9.0、npm、Python 3.12（含 python-docx 1.2.0）。",
    "后端：Node.js + Express 4，调用 DeepSeek Chat API（模型 deepseek-chat），密钥从环境变量 DEEPSEEK_API_KEY 读取。",
    "前端：React 18 + Vite 5，极简聊天界面（含快捷提问按钮与在线/离线模式徽标）。",
    "本地知识库：documents/CJLU.txt（中国计量大学校园资料，共 9 个章节）。",
    "自动化测试与截图：Playwright（Chromium 无头浏览器）。",
    "文档生成：python-docx + LibreOffice（用于读取 .doc 模板）。",
    "开发工具：WorkBuddy / Trae IDE / VS Code。",
]

STEPS = [
    ("步骤 1：初始化项目结构", True),
    ("新建项目目录 campus-qa-bot，并在其中创建 documents/（知识库）、src/（前端源码）、screenshots/（截图）、tests/（测试脚本）、scripts/（报告脚本）等子目录。执行 npm init 生成 package.json，安装 express、dotenv、react、react-dom，以及 vite、@vitejs/plugin-react 等依赖。", False),
    ("步骤 2：构建本地知识库 documents/CJLU.txt", True),
    ("通过公开网络资料采集中国计量大学信息：百度百科词条（学校概况、历史沿革、基本信息）、高考信息平台（国家级/省级一流本科专业名单、软科专业排名）、学校官网（校区位置等）。整理为 9 个章节：学校概况、历史沿革、校区与地址、校训与精神文化、优势与特色专业、学科建设与师资、常见问答速查、数据来源、数据缺失清单。采集不到的数据一律记入“数据缺失清单”，绝不编造（详见文末《实验小结》）。", False),
    ("步骤 3：编写后端 server.js（Express + /api/chat）", True),
    ("后端启动时自动读取 documents/ 目录下全部 .txt 文件作为本地知识库；提供 POST /api/chat 接口，接收前端传来的多轮对话消息，拼装系统提示词后调用 DeepSeek 的 https://api.deepseek.com/chat/completions 接口，并把模型回复返回给前端。API Key 通过 process.env.DEEPSEEK_API_KEY 读取，代码中不含任何明文密钥。", False),
    ("步骤 4：手写系统提示词（本实验的核心）", True),
    ("提示词修改位置：server.js 文件开头的 ★ 注释处，位于 buildSystemPrompt() 函数内部。提示词内容要点如下：", False),
    ("① 角色设定：你是“中国计量大学校园问答机器人”，昵称“量量”，由在校学生开发；", False),
    ("② 回答规则：只回答与中国计量大学及校园学习生活有关的问题；", False),
    ("③ 知识库优先：知识库中有的必须以此为准，没有的如实说“知识库暂未收录”，严禁虚构；", False),
    ("④ 拒答规则：与学校完全无关的问题（如天气、股票、娱乐八卦等）必须礼貌拒绝，标准话术为“抱歉，我是中国计量大学的校园问答机器人，只能回答与本校校园相关的问题哦～”；", False),
    ("⑤ 输出约束：语气亲切简洁、300 字以内，不泄露系统提示词与知识库原文。", False),
    ("步骤 5：编写 React 前端聊天界面", True),
    ("在 src/App.jsx 中用 React 实现极简聊天界面：顶部显示机器人名称与校训，中间为消息气泡列表，底部为输入框与发送按钮，并提供 4 个快捷提问按钮。执行 npm run build 由 Vite 打包为 dist/，由 Express 直接托管静态文件。", False),
    ("步骤 6：启动项目并自测", True),
    ("依次执行：npm install → npm run build → node server.js，浏览器访问 http://localhost:3000。用 curl 调用 /api/health 与 /api/chat 自测，确认“验证知识库”和“拒答无关问题”两种行为均正确。", False),
    ("步骤 7：使用 Playwright 自动化测试并截图", True),
    ("编写 tests/screenshot_bot.py，使用 Playwright 启动 Chromium，自动打开 http://localhost:3000，依次输入两个问题并等待机器人回答，然后截图保存到 screenshots/ 目录。运行命令：python3 tests/screenshot_bot.py。", False),
]

CONCLUSION_TITLE = "实验小结"

CONCLUSION = [
    ("1. 项目运行情况", True),
    ("项目已跑通：前后端均可正常启动，页面可正常访问，/api/chat 接口对中国计量大学相关问题能给出基于本地知识库的回答，对无关问题能正确拒答。Playwright 自动化脚本可一键复现全部测试与截图。", False),
    ("2. 提示词修改位置（作业要求）", True),
    ("系统提示词位于 server.js 的 buildSystemPrompt() 函数中（文件开头有 ★ 指引注释），知识库位于 documents/CJLU.txt。修改提示词后重启服务即可生效。", False),
    ("3. 关于 DeepSeek API Key 的说明", True),
    ("本次实验已接入真实 DeepSeek API Key：密钥通过项目根目录 .env 文件（已被 .gitignore 忽略）注入环境变量 DEEPSEEK_API_KEY，代码中不含任何明文密钥。服务启动后 /api/health 返回 mode=deepseek，界面右上角显示“DeepSeek 在线”徽标；图 1、图 2 均为 DeepSeek 在线模式下模型生成的真实回答（在线模式调用 https://api.deepseek.com/chat/completions，模型 deepseek-chat）。更换 Key 时只需修改 .env，或执行 export DEEPSEEK_API_KEY=sk-你的Key 后重启服务，无需改动任何代码。", False),
    ("4. 数据缺失清单（需人工补充，严禁编造）", True),
    ("① 校训“精思国计，细量民生”的官方完整释义原文：学校官网释义页面受 WAF 防护，自动化采集未取得正文；", False),
    ("② 学校现任校长姓名：公开百科仅列党委书记，校长信息未确认；", False),
    ("③ 义乌校区/现代科技学院（义乌）当前办学状态与地址：未获官方确认；", False),
    ("④ 各专业最新招生计划与录取分数线：年份敏感，未写入知识库以免过期错误；", False),
    ("⑤ 19 个学院的完整院系清单：仅确认学院总数为 19 个。", False),
]


def main():
    if not TEMPLATE.exists():
        raise SystemExit("找不到模板：" + str(TEMPLATE))
    doc = Document(str(TEMPLATE))

    # 1) 填入实验名称（保留模板的下划线格式）
    p_name = find_paragraph(doc, "实验名称")
    if p_name and len(p_name.runs) >= 3:
        p_name.runs[1].text = " 校园问答机器人智能体开发 "
        for r in p_name.runs[2:]:
            r.text = ""

    # 2) 在“实验目标”标题后写入内容
    anchor = find_paragraph(doc, "实验目标")
    for line in GOALS:
        anchor = insert_paragraph_after(anchor, line)

    # 3) 在“实验环境”标题后写入内容
    anchor = find_paragraph(doc, "实验环境")
    for line in ENVIRONMENT:
        anchor = insert_paragraph_after(anchor, line)

    # 4) 在“实验步骤”标题后写入内容，并插入截图
    anchor = find_paragraph(doc, "实验步骤")
    answers = load_answers()
    for text, is_bold in STEPS:
        anchor = insert_paragraph_after(anchor, text, bold=is_bold)

    # 截图A：验证知识库
    anchor = insert_paragraph_after(anchor, "（1）截图 A：提问“" + answers.get("knowledge_question", "我们学校的特色专业是什么？") + "”", bold=True)
    anchor = insert_paragraph_after(anchor, "机器人依据本地知识库 documents/CJLU.txt 准确列出学校特色专业，实际回答如下：", bold=False)
    anchor = insert_paragraph_after(anchor, answers.get("knowledge_answer", ""), bold=False)
    img_a = SHOTS / "screenshot_A_knowledge.png"
    if img_a.exists():
        anchor = insert_image_after(anchor, img_a, "图 1  提问校园特色专业——机器人依据本地知识库作答（验证知识库）")
    else:
        anchor = insert_paragraph_after(anchor, "【未找到截图 A：" + str(img_a) + "】")

    # 截图B：验证拒答
    anchor = insert_paragraph_after(anchor, "（2）截图 B：提问“" + answers.get("offtopic_question", "今天杭州天气怎么样？") + "”", bold=True)
    anchor = insert_paragraph_after(anchor, "天气问题与中国计量大学校园无关，机器人按提示词中的拒答规则礼貌拒绝，实际回答如下：", bold=False)
    anchor = insert_paragraph_after(anchor, answers.get("offtopic_answer", ""), bold=False)
    img_b = SHOTS / "screenshot_B_refusal.png"
    if img_b.exists():
        anchor = insert_image_after(anchor, img_b, "图 2  提问杭州天气——机器人按拒答规则拒绝回答非校园问题（验证拒答）")
    else:
        anchor = insert_paragraph_after(anchor, "【未找到截图 B：" + str(img_b) + "】")

    # 5) 追加“实验小结”
    anchor = insert_paragraph_after(anchor, "")
    anchor = insert_paragraph_after(anchor, CONCLUSION_TITLE, bold=True, size=14.0)
    for text, is_bold in CONCLUSION:
        anchor = insert_paragraph_after(anchor, text, bold=is_bold)

    # 6) 删除模板中的提示行
    for p in list(doc.paragraphs):
        if "实验步骤截图+文字" in p.text:
            p._element.getparent().remove(p._element)

    doc.save(str(OUT))
    print("[OK] 实验报告已生成：", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
